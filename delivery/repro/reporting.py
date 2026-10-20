"""Create an honest result report from persisted outputs, without invented claims."""
from pathlib import Path
import json
import os
import shutil
import subprocess
import tempfile

def style_document(doc):
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    section=doc.sections[0]
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=Inches(.7);section.bottom_margin=Inches(.7)
    section.left_margin=Inches(.75);section.right_margin=Inches(.75)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
        style=doc.styles[name]
        style.font.name='Noto Sans CJK SC'
        style.font.color.rgb=RGBColor(0,0,0)
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
        for attr in ('asciiTheme','eastAsiaTheme','hAnsiTheme','cstheme'):
            style._element.get_or_add_rPr().rFonts.attrib.pop(qn('w:'+attr),None)
        pr=style._element.find(qn('w:pPr'))
        if pr is not None:
            for child in list(pr):
                if child.tag in (qn('w:pBdr'),qn('w:numPr')):pr.remove(child)
    doc.styles['Subtitle'].font.italic=False
    doc.styles['Normal'].font.size=Pt(10)
    doc.styles['Normal'].paragraph_format.line_spacing=1.15
    doc.styles['Normal'].paragraph_format.space_after=Pt(5)
    doc.styles['Title'].font.size=Pt(23)
    doc.styles['Heading 1'].font.size=Pt(15)
    doc.styles['Heading 2'].font.size=Pt(12)
    doc.styles['Heading 1'].paragraph_format.space_before=Pt(10)
    doc.styles['Heading 1'].paragraph_format.space_after=Pt(5)
    footer=section.footer.paragraphs[0]
    footer.alignment=2
    run=footer.add_run('MolRepBench  ')
    run.font.size=Pt(8)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

def table(doc, headers, rows, widths=None):
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    t=doc.add_table(rows=1, cols=len(headers));t.autofit=False
    if widths:
        for col,width in zip(t.columns,widths):col.width=Inches(width)
    for cell,text in zip(t.rows[0].cells,headers):cell.text=str(text)
    for row in rows:
        for cell,text in zip(t.add_row().cells,row):cell.text=str(text)
    for idx,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit');trPr.append(cant)
        if idx==0:
            repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
        for col,cell in enumerate(row.cells):
            if widths:cell.width=Inches(widths[col])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for side in ('top','left','bottom','right'):
                b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'5');b.set(qn('w:color'),'D9D9D9');borders.append(b)
            pr.append(borders)
            margins=OxmlElement('w:tcMar')
            for side in ('top','left','bottom','right'):
                m=OxmlElement('w:'+side);m.set(qn('w:w'),'90');m.set(qn('w:type'),'dxa');margins.append(m)
            pr.append(margins)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DDE8F1' if idx==0 else ('F6F7F8' if idx%2==0 else 'FFFFFF'));pr.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.15
                for r in p.runs:r.font.size=Pt(9);r.bold=idx==0
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

def report(run_dir, output):
    from docx import Document
    from .core import ROOT, read_jsonl
    run_dir, output=Path(run_dir),Path(output)
    output.mkdir(parents=True,exist_ok=True)
    status_path=run_dir/'evaluation_status.json'
    if not status_path.exists():
        raise ValueError('Run score on actual saved responses first. No success report is created without evidence.')
    status=json.loads(status_path.read_text(encoding='utf-8'))
    summary=json.loads((run_dir/'summary.json').read_text(encoding='utf-8'))
    metadata=json.loads((run_dir/'run_metadata.json').read_text(encoding='utf-8')) if (run_dir/'run_metadata.json').exists() else {}
    doc=Document();style_document(doc)
    doc.add_paragraph('分子文本表示实验结果记录',style='Title')
    doc.add_paragraph('依据保存的真实响应与评分文件生成 需人工核对后汇报',style='Subtitle')
    state='计划中的模型调用已全部保存' if status['inference_complete'] else '实验尚未完成'
    paragraphs=[f'当前状态为{state}。预定 {status["expected_requests"]} 条请求，API 成功 {status["api_success"]} 条，能解析最终答案 {status["parsed_answers"]} 条。调用成功不代表答对，也不代表已完整复现论文。',
      '本次执行的是基于作者公开数据、经审计后适配的复现协议。B5 标签规则、分子对成员表示、选项顺序、生成样本数及停止序列等与公开脚本存在记录的差异。论文与公开 B9/B10 数据源说明也不一致。',
      '原论文研究同一个分子在不同文本表示下的 LLM 表现是否改变。RDKit 用于结构解析、格式转换及标准答案计算；模型只获得任务提示词，标准答案不发送给模型。']
    for p in paragraphs:doc.add_paragraph(p)
    doc.add_heading('模型与证据',level=1)
    cfg=metadata.get('settings',{})
    for line in [f'结果目录 {run_dir.resolve()}',f'结果类型 {metadata.get("result_kind","未记录")}',
       f'模型 {cfg.get("model","外部响应 未核实模型身份")}',f'推理标签 {cfg.get("thinking","未记录")} 实际开关须核对后端',
       '原始回答 raw.jsonl  逐题评分 scored.jsonl  汇总 summary.csv  状态 evaluation_status.json']:
        doc.add_paragraph(line)
    doc.add_heading('分任务与分表示成绩',level=1)
    doc.add_paragraph('空值表示未完成、无法计算或解码未评估。连续性质仅在成功解析的数字上计算相关性，同时显示覆盖率。B2 缺答案时主 Macro-F1 留空。')
    rows=[]
    for r in summary:
        key=('macro_f1' if r['task']==2 else 'spearman_rho' if r['task']==3 and r['subtask'] in ('logp','tpsa')
             else 'exact_match' if r['task']==6 else 'accuracy_all')
        value=r.get(key)
        rows.append([f'B{r["task"]} {r["subtask"]}',r['representation'],r['n_expected'],
                     f'{r["parse_coverage"]:.1%}',f'{value:.4f}' if value is not None else '未评估'])
    table(doc,['任务','表示','计划数','解析率','主指标'],rows,[1.35,2.1,.6,.8,1.65])
    doc.add_heading('可逐条追溯的案例',level=1)
    scored=read_jsonl(run_dir/'scored.jsonl')
    raw={r['id']:r for r in read_jsonl(run_dir/'raw.jsonl')}
    manifest_path=metadata.get('manifest_path')
    manifest={r['id']:r for r in read_jsonl(Path(manifest_path)/'manifest.jsonl')} if manifest_path else {}
    cases=[r for r in scored if not r.get('correct') and r['api_status']=='ok'][:3]
    if not cases:doc.add_paragraph('当前保存结果中没有可自动选出的真实错误案例。需人工查看原始回答；不生成虚构案例。')
    for idx,r in enumerate(cases,1):
        doc.add_heading(f'案例 {idx}',level=2)
        doc.add_paragraph(r['id'])
        if r['id'] in manifest:doc.add_paragraph('实际 prompt\n'+manifest[r['id']]['prompt'])
        doc.add_paragraph('模型原始回答\n'+raw.get(r['id'],{}).get('raw_response','未保存'))
        doc.add_paragraph('参考答案 '+str(r['ground_truth'])+'  解析结果 '+str(r['predicted']))
    doc.add_heading('我可以据此解释的内容',level=1)
    doc.add_paragraph('我比较的是模型读取不同分子写法后的答题表现。输入格式的显式程度、语言线索和可保留的信息不同，因此需要同时检查任务、模型、样本和格式转换。仅从一张成绩表不能证明某种表示对所有任务最好。')
    doc.add_paragraph('汇报前我还需确认 IUPAC 输入审核、生成解码、FCD、配对置信区间及模型版本；原文裁判图和机制分析需另列完成情况。程序未自动写入任何“全文成功复现”的结论。')
    doc.add_heading('尚需人工完成',level=1)
    doc.add_paragraph('逐项对照论文结果；解释真实失败案例；把实际模型、数据与协议偏差补齐；确认哪些结论被本次实验支持。请同时阅读随包的论文理解与复现指南。')
    docx=output/'experiment_results.docx';doc.save(docx)
    pdf=output/'experiment_results.pdf'
    # Prefer one rendering for both artifacts; fall back to a transparent text-layout PDF.
    runtime_office=Path(os.environ.get('CODEX_PRIMARY_RUNTIME_ROOT','/nonexistent'))/'dependencies/bin/override/soffice'
    office=str(runtime_office) if runtime_office.exists() else (shutil.which('libreoffice') or shutil.which('soffice'))
    if office:
        with tempfile.TemporaryDirectory() as temp:
            result=subprocess.run([office,'-env:UserInstallation=file://'+temp+'/profile','--headless',
                '--convert-to','pdf','--outdir',str(output.resolve()),str(docx.resolve())],capture_output=True,text=True,timeout=180)
        if result.returncode!=0 or not pdf.exists():
            raise RuntimeError('DOCX saved, but PDF rendering failed. Inspect office logs and render before delivery.')
    else:
        from reportlab.pdfgen import canvas
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.lib.utils import simpleSplit
        pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
        c=canvas.Canvas(str(pdf),pagesize=(612,792));y=750
        text=['分子文本表示实验结果记录','PDF 使用文本布局导出 与 DOCX 分页可能不同']
        text += [p.text for p in doc.paragraphs if p.text]
        text += [' | '.join(str(v) for v in r) for r in rows]
        for block in text:
            for line in simpleSplit(block,'STSong-Light',10,512):
                if y<45:c.showPage();y=750
                c.setFont('STSong-Light',10);c.drawString(50,y,line);y-=15
            y-=5
        c.save()
    return {'docx':str(docx),'pdf':str(pdf),'status':state,'requires_human_review':True}
