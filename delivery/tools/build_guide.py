"""Build the initial Chinese reading guide; this is not a completed LLM report."""
from pathlib import Path
import json
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from repro.core import ROOT, read_jsonl
from repro.reporting import style_document, table
from docx import Document
from docx.shared import Pt

def main():
    doc=Document();style_document(doc)
    def p(text):doc.add_paragraph(text)
    def h(text):doc.add_heading(text,level=1)
    def sub(text):doc.add_heading(text,level=2)
    def page(title):
        paragraph=doc.add_heading(title,level=1)
        paragraph.paragraph_format.page_break_before=True
    def code(text):
        para=doc.add_paragraph()
        para.paragraph_format.line_spacing=1.05
        para.paragraph_format.space_after=Pt(6)
        r=para.add_run(text);r.font.name='DejaVu Sans Mono';r.font.size=Pt(8.5)
    plan=json.loads((ROOT/'plans/paper/plan.json').read_text())
    audit=json.loads((ROOT/'evidence/dataset_audit.json').read_text())
    smoke=read_jsonl(ROOT/'plans/smoke/manifest.jsonl')
    b1=next(r for r in smoke if r['task']==1 and r['representation']=='canonical_smiles')

    doc.add_paragraph('分子文本表示论文理解与复现指南',style='Title')
    doc.add_paragraph('Rethinking Molecular Text Representations for LLMs\n供教授汇报准备与 CodeBuddy 本地执行使用',style='Subtitle')
    p('本项目已找到并取得作者公开代码与数据，完成便携 Python 实现、离线数据审计和软件检查。目前没有执行真实 LLM 推理，因此本文件是论文理解与运行指南，不能作为“已成功重现论文成绩”的结果报告。')
    h('这次要复现什么')
    p('论文的问题是：同一个分子换一种文本写法，大语言模型是否仍能理解并完成相同任务？作者比较 9 种分子文本表示、8 类任务和 16 个 LLM 配置。研究对象是模型对不同表示的适应能力，不是在这里训练一个新的分子神经网络。')
    p('最小实验单元是一道固定问题。把同一分子分别写成 SMILES、IUPAC、CML 等格式，保持题目和模型条件一致，保存实际回答，再与 RDKit 算出的标准答案比较。表示改变后成绩若改变，才得到可检查的实验现象。')
    h('已完成与待完成')
    table(doc,['内容','本次状态'],[
        ['官方来源','作者 arunraja-hub/molrepbench；固定 commit 21d36941cddf0dc1ffb5f3590296592dc68e4ac8'],
        ['可运行流程','已实现 8 类任务的 9 种表示输入、提示词、推理接口、评分与结果报告入口'],
        ['离线验证','已生成 270 条 smoke 与 29139 条 paper 请求清单，已完成数据审计与软件测试'],
        ['真实模型成绩','未运行；需提供本地模型服务或已有 API 配置'],
        ['全文所有图表','未完成；FCD、原裁判与机制分析有单独依赖或待核实细节']],[1.25,5.75])
    p('原始论文为用户提供的 arXiv:2606.03057v1，2026 年 6 月 2 日。来源与本项目差异详见第 8 页及 README。软件测试使用的模拟回答只验证程序行为，不是 LLM 化学能力。')

    page('九种表示分别在写什么')
    p('分子可看作一张图：原子是节点，化学键是边。表示就是把这张图写成模型能读取的文本。格式还可能携带立体、同位素、电荷等信息；并非每个转换实现都保存全部信息。')
    table(doc,['表示','为何纳入论文比较','阅读与复现时注意'],[
        ['Canonical SMILES','常用紧凑基线，以规范顺序遍历原子','作者转换显式关闭立体信息；“canonical”本身不等于不含立体'],
        ['Isomeric SMILES','在 SMILES 中加入立体及同位素标记','看能否利用 @、斜线等标记区分立体异构体'],
        ['Randomized SMILES','同一图使用不同遍历顺序','检验是否依赖固定字符串；本包按 seed 冻结具体输入'],
        ['DeepSMILES','减少配对括号和环闭合编号的语法困难','语法便利不等于理解更好；作者从非立体 SMILES 转换'],
        ['IUPAC name','名称中的化学词根与自然语言描述有联系','命名或标准化可能折叠互变形式/电荷态，须验证成员映射'],
        ['SELFIES','设计上提高合法分子解码的稳健性','合法分子未必满足描述；实际输入仍须完整合法且可解码'],
        ['MolJSON','用 JSON 显式列出原子和化学键','作者所用 schema 不支持立体；不能推广为所有 JSON 格式都如此'],
        ['CML','用 XML 显式列出原子、键及属性','更长；转换器与具体字段决定保留哪些化学信息'],
        ['InChI','分层写出分子式、连接和状态等信息','标准 InChI 可能规范化某些互变形式，不能仅按字符串不同判错']],[1.35,2.7,2.95])
    sub('为何不是随便选九种')
    p('论文 §3 明确按现代化学信息工具和深度学习中的常用程度选择。它们覆盖常用字符串、为生成而改写的语法、具有语言线索的名称和显式图结构。比较这些选择，可以检验“更常见”“更容易生成”或“图信息更显式”是否真的适合不同任务。')
    p('CML/MolJSON 在结构题上可能更容易读取，但更长、成本更高；IUPAC 可能帮助语义对齐，但前提是名称对应正确结构。这里的“可能”是解释假设，不能代替本次实际结果。')

    page('八类任务分别测什么')
    table(doc,['任务','要模型做什么','论文规模与主指标'],[
        ['B1 原子计数','数指定元素 C N O S F Cl 的原子数','200 分子；每分子 1 元素；exact accuracy'],
        ['B2 官能团','分别判断伯胺、酯、醛、磺酰胺、卤代基是否存在','200 × 5 问；Macro-F1'],
        ['B3 性质估计','估计 LogP、TPSA、HBD 和 HBA','200 × 4 问；连续性质 Spearman，计数 exact accuracy'],
        ['B4 分子检索','根据自然语言描述，从四个结构中选一个','78 问；top-1 accuracy'],
        ['B5 同分子判断','判断两种写法是否代表完全相同的分子','500 对；accuracy；立体异构体为 No'],
        ['B6 描述生成分子','根据描述生成指定表示的分子','200 描述；有效性、exact match、Tanimoto、FCD'],
        ['B9 互变异构体','判断两结构是否属于同一规范互变异构体','250 对；accuracy'],
        ['B10 质子化态','判断中和后是否是同一结构','210 对；accuracy']],[1.15,3.0,2.85])
    sub('为何需要不同层次的题')
    p('原子计数检验能否读出基本组成；官能团要求识别局部连接关系；性质估计要求从图总结化学特征。检索把结构与自然语言联系起来，生成则反向把描述变成结构。三类成对任务分别考查身份、互变关系与质子化关系，不能混成一个“相似度”问题。')
    sub('为什么不直接选药效或毒性')
    p('论文 §2 的理由是，很多生化性质依赖实验测量及化学或物理归纳偏置。用这些性质评估文本表示，容易混入模型是否掌握相关实验规律的因素。本文选择能由 RDKit 从图计算的 LogP、TPSA、HBD/HBA，让标准答案更可重复。RDKit 算法值也不等于真实实验测量。')
    p('B2 的五类标签用于避免过于常见、容易被简单文本线索识别的官能团占主导；具体选择理由见作者 config.py 注释。SMARTS 定义也是实验的一部分，例如作者 halide 模式限定与 sp3 碳连接的卤素，不能自行扩展含义。')

    page('零基础需要分清的化学和指标')
    sub('同一分子不等于同一字符串')
    p('例如 CCO 和 OCC 都可写乙醇。它们字符串不同，连接图相同，B5 应回答 Yes。立体异构体连接关系可能相同，却具有不同三维排列，B5 应回答 No。互变异构体涉及氢位置和键重排，B9 才问它们能否归为同一互变家族。酸与去质子化后的阴离子属于 B10 的另一种关系。')
    sub('四个性质是什么')
    p('LogP 描述分子在两种相中的分配倾向，这里采用 RDKit 的 Wildman-Crippen 算法估计。TPSA 是拓扑极性表面积。HBD 是能作为氢键供体的位点数，HBA 是氢键受体位点数；它们不是简单数所有氢、氮或氧。')
    table(doc,['指标','直观解释','为什么这样用'],[
        ['准确率','答对题数除以计划题数','原子数和离散判断有明确标准答案'],
        ['Macro-F1','先算五种官能团各自 F1，再平均','让较少出现的官能团也有相同权重；不是总准确率'],
        ['Spearman 相关','比较预测与真值的排序是否一致','适合 LogP/TPSA 的相对排序；不能说明绝对误差很小'],
        ['MAE','平均偏差的绝对值','本包作为连续性质的补充指标，不替换论文主指标'],
        ['Validity','生成串能否解析成合法分子','只检验可解析，不检验是否符合描述'],
        ['Exact match','生成图是否与参考图相同','主列沿用作者忽略立体的实现，另报含立体列'],
        ['Tanimoto','比较分子指纹中共同的结构特征','相似不等于完全正确；须说明指纹与无效输出处理'],
        ['FCD','比较生成集合与参考集合的 ChemNet 分布','是集合指标，不能评价某一个分子，少量样本不稳定']],[1.2,2.65,3.15])
    p('论文正文 §3 与表 2 为 LogP/TPSA 指定 Spearman；附录 D 的“exact match”措辞存在不一致。本包按正文及公开评分代码的连续性质处理，并显式记录这一选择。QED 不属于本论文 B3。')

    page('核心代码如何串成一个实验')
    p('以 B1 为例，先读入原始 SMILES，让 RDKit 构建分子图并计算参考答案，再把输入换成各表示。LLM 只看任务和该表示文本；标准答案留在本地，等模型回答后评分。')
    code('counts = chemistry.get_atom_counts(mol)\nfor rep in reps:\n    s = enc(row, rep, mid)\n    add(1, mid, element, rep,\n        prompts.create_atom_counting_prompt(s, element, rep),\n        counts[element])')
    p('上面节选自 repro/core.py 的实际实现；对失败转换另有 excluded 记录。add 将 ID、表示、实际 prompt、参考答案和 SHA256 写进 manifest。hash 用于确认不同模型是否回答了同一道输入。')
    code("body = {\n  'model': cfg['model'],\n  'messages': [{'role': 'user', 'content': item['prompt']}],\n  cfg.get('token_parameter', 'max_tokens'):\n      cfg.get('max_tokens', 2048)\n}")
    p('这段对应 repro/inference.py 的请求构建逻辑。ground_truth 没有放入 messages。若直接把 RDKit 算出的答案当成“模型结果”，就没有测试 LLM；如果把答案放进 prompt，又会造成答案泄漏。')
    table(doc,['文件','读它能回答什么'],[
        ['repro/core.py','样本怎么选、9 种表示怎么准备、每道题怎么固定'],
        ['vendor 的 utils/prompts.py','作者原始任务提示词是什么'],
        ['vendor 的 config.py 和 chemistry.py','官能团 SMARTS、元素、性质的标准答案怎么定义'],
        ['vendor 的 representations.py 和 moljson/conversion.py','格式转换保存或丢失了哪些信息'],
        ['repro/inference.py','模型实际收到什么、如何保存原回答及续跑'],
        ['repro/scoring.py','如何解析最后答案、计分并保留失败分母']],[2.65,4.35])
    p('本包的新增工作是便携执行、审计和可追溯评分。作者已有的格式与提示词实现保留来源；不能把整个 vendor 目录称为自己原创。')

    page('实际任务提示词如何使用')
    p('任务 prompt 沿用作者源码。除输入的分子表示变化外，任务文字保持一致。模型推理开关通过后端配置控制，不能简单删掉“reason step by step”就声称对应论文的 non-reasoning 条件。')
    sub('一条已经生成的真实 B1 请求')
    code(b1['prompt'])
    p(f'问题 ID 为 {b1["id"]}；RDKit 参考答案为 {b1["ground_truth"]}。这是请求清单中的真实问题；当前尚无模型回答，因此不能给出模型得分。')
    sub('其他任务的提问形式')
    table(doc,['任务','作者 prompt 的问题部分'],[
        ['B2','Does the following molecule contain a {functional_group}?'],
        ['B3','Estimate the {property_display} of the following molecule.'],
        ['B4','Which of the following molecules matches this description? 后接描述及 A 到 D 四个表示'],
        ['B5','Do the following two molecular representations refer to the same molecule?'],
        ['B9','Are the following two molecules tautomeric forms of the same compound?'],
        ['B10','Are the following two molecules different protonation states of the same compound?']],[.65,6.35])
    p('这些题要求把最后数字、Yes/No 或选项字母放入 LaTeX boxed 标记。本包只读取明确的最终答案或纯答案，不把推理正文中随意出现的第一个数字当作答案。')
    p('docs/actual_prompt_examples.txt 收录 8 类任务各一条完整 canonical SMILES 实际请求及其参考答案。全部 9 种表示的实际请求在重新运行 prepare 后的 manifest.jsonl。')

    page('生成任务与运行步骤')
    sub('B6 与其他任务有什么不同')
    p('B6 的输入是 ChEBI 的自然语言描述，输出才是指定表示的分子。不能把参考 SMILES 作为生成模型的输入。作者额外要求简短回答，避免复述描述；本包沿用同一 generation prompt builder，要求 JSON 格式输出。')
    code('Given a description of a molecule, generate\nthe corresponding {rep_name} string.\n\nDescription: {description}\n\nFor non-MolJSON output:\n{"molecule": "..."}')
    p('MolJSON 则直接输出含 atoms、bonds、charges、aromatic_n_h 的对象，并按作者 schema 约束。其他格式即使包在 JSON 中，内部字符串仍必须是对应分子表示。schema 只约束形状，不能保证化学含义正确。')
    sub('本地执行顺序')
    code('python reproduce.py doctor\npython reproduce.py verify-vendor\npython reproduce.py setup\npython -m unittest discover -s tests -v\npython reproduce.py prepare --profile smoke --out plans/smoke\npython reproduce.py run --plan plans/smoke \\\n  --settings settings.json --out results/smoke --max-requests 1')
    p('上述 python 应为项目 .venv 中的解释器。模型服务与 settings 准备好后先调用一条；确认真实响应、token 用量和最后答案，再续跑 270 条 smoke。smoke 只检查链路，不能据此选“最佳表示”。')
    sub('转入正式比较')
    p('先审核 IUPAC 输入，再生成 paper_verified 清单。明确模型、模式与预算后运行正式实验，执行 score、bootstrap 和 report。模型、参数、输入或 seed 改变后建立新结果目录，保留旧结果，不能混在同一份 JSONL 中。')
    p('本次公开名称路线共 29139 请求/模型/条件，其中 B5 少 1 个 IUPAC 输入，B10 少 2 个 IUPAC 输入。已审核路线的数量可能更小，需按实际 plan 与共同样本报告。')

    page('公开材料的差异与本包处理')
    p('可运行不代表与作者当年实验完全相同。以下差异已经检查，必须进入向教授的汇报，不能把它们当作无关的安装问题。')
    table(doc,['证据','当前处理与含义'],[
        ['B5 的 375 对标签差异','CSV 250 Yes / 250 No 与 RDKit 身份判断一致；本包使用 CSV。作者 runner/evaluator 的 natural/stereoisomer 规则不符合当前 CSV 与同分子问题'],
        ['B9/B10 的数据源','公开 CSV 全标 ChEBI-20，论文附录 C 写 ZINC250K。保留公开数据来源声明，未重建未核实的 ZINC 实验'],
        ['预存 SELFIES 成员不匹配','B9 有 183 个、B10 有 118 个成员解码后不等于对应 SMILES。本包非 IUPAC 表示全部从每个成员结构重新生成'],
        ['互变异构体标签争议','当前 RDKit 默认规则下有 4/250 对不一致，枚举也出现默认限制警告。保留公开标签并记录索引；原因待版本和算法条件核实'],
        ['IUPAC 名称可能折叠状态','B9 125 对、B10 105 对不同结构共享同名。须借 OPSIN 审核，区分信息折叠与模型错误'],
        ['解码与采样差异','B6 按论文取 200，公开配置为 250；本包固定随机输入与共同选项顺序，且取消空行停止。都记录为协议改动']],[1.7,5.3])
    sub('如何理解 IUPAC 的审核路线')
    p('released 路线保留公开名称，并标为输入未验证。verified 路线只保留 OPSIN 解码后与成员 SMILES 的立体身份完全一致的名称。它不能自动修复坏映射；不一致、网络失败和无法解析各有记录。过滤后必须比较共同样本，不能只比较聚合排名。')
    p('“同一个 IUPAC 名称出现在两个状态”本身不必然证明名称错误，但它意味着模型无法只凭该字符串区分这些状态。是否应标准化、是否应保留某种差异，是任务定义和表示能力的问题。')

    page('怎样解释未来得到的成绩')
    sub('先看证据完整性')
    p('查看预定条数、实际调用成功数、最终答案解析率及解码状态。API 成功返回并不保证有最终答案；finish_reason 为 length 表示可能截断。IUPAC 未联网解码时应标未评估，不记为无效分子。')
    sub('再看公平比较')
    p('比较表示时使用同一模型、相同问题和可比较的推理/解码条件。不同表示有缺失时，应在共同问题 ID 上做配对比较，并同时报告完整计划覆盖率。不同模型或样本规模的两行成绩，不能直接用来证明表示带来的差异。')
    p('连续性质仅在可解析数字上计算 Spearman 和 MAE，须同时显示覆盖率。官能团若有无法解析的回答，本包不填主 Macro-F1，只给 valid-only 值作参考及全体准确率。离散题的缺答案按计划分母计错。')
    sub('论文结论与我的结果要分开')
    p('原论文报告 CML/MolJSON 更适合某些结构任务、IUPAC 更适合语义任务，并强调没有一种表示在所有任务上最好。这是论文作者在其模型和设置下的观察。我的正式实验只有实际跑完并保存证据后，才能说明是否支持同样的趋势。')
    p('若趋势不同，先检查标签、数据源、格式转换、模型版本、推理模式、token 预算、schema、样本缺失和评分定义。排除实现差异后，再讨论模型本身的变化。不能直接把差异解释为我改进了原论文。')
    sub('哪些内容仍然属于扩展范围')
    table(doc,['部分','完成它需要什么'],[
        ['FCD','真实生成结果、ChemNet 依赖和权重、足够样本；本次未计算'],
        ['LLM 裁判','对齐论文 Gemini 3 Flash、样本与 rubric；本包 judge 是标注清楚的改编入口'],
        ['Tokenizer 审计','实际 tokenizer 与固定分子样本；提供入口但尚未下载/运行'],
        ['线性探针与注意力','核实层、span、池化、训练划分及 GPU；公开文件未提供可直接重放全部图的完整流程']],[1.6,5.4])
    p('这些项目不应因为主任务脚本跑通而自动标为成功。项目状态文件始终把全文复现状态与主任务执行分开。')

    page('向教授汇报时怎样说清楚')
    sub('目前可以如实汇报的版本')
    p('我已找到原论文的作者仓库，核对了九种分子表示、八类任务及实际提示词。我把作者代码和公开数据固定到一个 commit，并实现了可以在本地 Python 环境运行的流程。现在完成的是数据、提示词和软件验证，还没有运行真实模型，所以还不能报告 LLM 复现成绩。')
    p('我理解这篇论文是在测试“分子的写法会不会改变模型的能力”。我会固定分子、题目和模型，改变表示，再保存原始回答并与 RDKit 标准答案对比。原子计数测基本解析，官能团与性质估计测结构推理，检索和生成测语言与结构的联系，成对任务则区分身份、互变异构体和质子化关系。')
    p('审读代码后我发现，公开脚本的 B5 标签读取与数据不一致，部分分子对预存表示也需要重新生成。公开 B9/B10 的数据源说明与论文有差异。我会明确记录这些修正，因此当前计划是公开材料的审计复现，不先声称逐个数值重现作者表格。')
    sub('运行成功后再补充的内容')
    p('实际报告中补入精确模型版本、运行条件、完成条数、每种表示成绩、缺失分母、置信区间和真实失败案例。口述时选一道真实题，展示它的 SMILES、变换后的输入、实际 prompt、模型原始回答、标准答案与评分，逐步解释每一项。不要在这些证据产生前写“我成功验证了某表示最好”。')
    sub('教授可能追问的具体问题')
    p('问 为什么选这些任务？\n答 任务覆盖组成、局部图、图属性、语言对齐与不同化学等价关系；标准答案大多可以计算，便于把差异追溯到表示和模型。')
    p('问 你的核心代码和 prompt 在哪？\n答 prepare 生成清单，作者 prompts.py 决定任务措辞，inference.py 只发送 prompt，scoring.py 按保存的原始回答重新计分。可现场打开同一 ID 的完整链路。')
    p('问 这是否已经完整复现？\n答 当前没有真实 LLM 结果；后续也必须分别说明主任务、模型覆盖、FCD、裁判和机制研究的完成状态。数据与协议偏差会保留，不用“全部完成”掩盖。')
    p('来源索引 论文 §2 §3 §4 §6 附录 C D E H I J；作者仓库固定 commit；README；UPSTREAM_PROVENANCE.json；evidence/dataset_audit.json。联网来源 https://arxiv.org/abs/2606.03057 和 https://github.com/arunraja-hub/molrepbench。')

    out=ROOT/'docs'/'paper_understanding_and_reproduction_guide.docx'
    doc.save(out)
    examples=[]
    for task in [1,2,3,4,5,6,9,10]:
        row=next(r for r in smoke if r['task']==task and r['representation']=='canonical_smiles')
        examples.append(f'=== B{task} {row["id"]} ===\nPROMPT\n{row["prompt"]}\n\nREFERENCE ANSWER (not sent to model)\n{row["ground_truth"]}\n\nMODEL RESPONSE\nNot run\n')
    (ROOT/'docs/actual_prompt_examples.txt').write_text('\n\n'.join(examples),encoding='utf-8')
    print(out)

if __name__=='__main__':main()
